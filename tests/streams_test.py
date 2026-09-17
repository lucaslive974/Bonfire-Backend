from io import BytesIO
from unittest.mock import MagicMock

import numpy as np
import pandas as pd
import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph

from domain.entities import AutoInfracao
from infrastructure.parsers.exceptions import (
    DocumentReadError,
    IncorrectInstanceError,
    InvalidDocumentDataError,
    PublicationDateNotFoundError,
    QuantityOfAtasMismatchError,
)
from infrastructure.parsers.pyingestion.streams import (
    BonfireInfracaoWriteStream,
    BonfireRecursoWriteStream,
    InfracoesCsvInputStream,
    InfracoesTransformStream,
    InfracoesXlsInputStream,
    RecursosDocxInputStream,
    SanitizedTextIO,
)

# ==========================================
# WRITE STREAMS TESTS
# ==========================================


def test_bonfire_recurso_write_stream_primeira_instancia():
    mock_db_manager = MagicMock()
    mock_session = mock_db_manager.session.return_value.__enter__.return_value
    mock_repo = mock_session.get_recurso_repository.return_value
    mock_repo.insert_primeira_instancia.return_value = 2

    stream = BonfireRecursoWriteStream(
        mock_db_manager, first_instance=True, batch_size=2
    )

    # Write a single item and a list of items
    stream.write({"NUM_RECURSO": "1"})
    stream.write([{"NUM_RECURSO": "2"}])
    stream.flush()

    assert mock_repo.insert_primeira_instancia.call_count >= 1


def test_bonfire_recurso_write_stream_segunda_instancia():
    mock_db_manager = MagicMock()
    mock_session = mock_db_manager.session.return_value.__enter__.return_value
    mock_repo = mock_session.get_recurso_repository.return_value
    mock_repo.insert_segunda_instancia.return_value = 1

    stream = BonfireRecursoWriteStream(
        mock_db_manager, first_instance=False, batch_size=10
    )
    stream.write({"NUM_RECURSO": "10"})
    stream.flush()

    mock_repo.insert_segunda_instancia.assert_called_once()


def test_bonfire_infracao_write_stream():
    mock_db_manager = MagicMock()
    mock_session = mock_db_manager.session.return_value.__enter__.return_value
    mock_repo = mock_session.get_autoinfracao_repository.return_value
    mock_repo.insert_bulk.return_value = 1

    stream = BonfireInfracaoWriteStream(mock_db_manager, batch_size=5)
    stream.write(
        {
            "NUM_AI": "12345-A",
            "NUM_NOTF": "NOTF-12345",
            "TIP_PENL": "MULTA",
            "NOM_CONC": "Consórcio BH Leste",
            "COD_LINH": "61",
            "NOM_LINH": "Estação Vilarinho",
            "DAT_OCOR_INFR": "2026-08-28T09:30:00",
            "COD_IRRG_FISC": 101,
            "ARTIGO": "Art. 1",
            "QTE_PONT": 3,
            "DAT_EMIS_NOTF": "2026-08-28T10:00:00",
            "DAT_LIMT_RECU": "2026-09-28T10:00:00",
            "VAL_INFR": 150.50,
        }
    )
    stream.flush()

    mock_repo.insert_bulk.assert_called_once()
    called_arg = mock_repo.insert_bulk.call_args[0][0]
    assert len(called_arg) == 1
    assert isinstance(called_arg[0], AutoInfracao)
    assert called_arg[0].notice_number == "12345-A"


# ==========================================
# SANITIZED TEXT IO TESTS
# ==========================================


def test_sanitized_text_io_accents_and_empty():
    empty_stream = BytesIO(b"")
    sanitized_empty = SanitizedTextIO(empty_stream)
    assert sanitized_empty.read() == ""

    # Stream with Portuguese accents and latin-1 chars
    raw_data = "Ação de Trânsito & Concessão".encode("utf-8")
    binary_stream = BytesIO(raw_data)
    sanitized = SanitizedTextIO(binary_stream)

    result = sanitized.read()
    assert result == "Acao de Transito & Concessao"

    sanitized.seek(0)
    assert binary_stream.tell() == 0


# ==========================================
# CSV & XLS INPUT STREAM TESTS
# ==========================================


def test_infracoes_csv_input_stream_detect_separator():
    stream = InfracoesCsvInputStream()

    pipe_data = BytesIO(b"COL1|COL2|COL3\nVAL1|VAL2|VAL3")
    assert stream._detect_separator(pipe_data) == "|"

    comma_data = BytesIO(b"COL1,COL2,COL3\nVAL1,VAL2,VAL3")
    assert stream._detect_separator(comma_data) == ","

    # Broken stream fallback
    broken_stream = MagicMock()
    broken_stream.read.side_effect = Exception("Read error")
    assert stream._detect_separator(broken_stream) == ";"


def test_infracoes_csv_input_stream_read():
    csv_bytes = b"NUM_AI;DAT_OCOR_INFR\n12345-A;01/01/2026\n"
    stream = InfracoesCsvInputStream()

    chunks = list(stream.read(BytesIO(csv_bytes)))
    assert len(chunks) == 1
    assert "NUM_AI" in chunks[0].columns

    # Test error wrapping
    bad_stream = MagicMock()
    bad_stream.read.side_effect = RuntimeError("Fatal disk error")
    with pytest.raises(DocumentReadError):
        list(stream.read(bad_stream))


def test_infracoes_xls_input_stream_error():
    stream = InfracoesXlsInputStream()
    with pytest.raises(DocumentReadError):
        list(stream.read(BytesIO(b"not an excel file")))


# ==========================================
# TRANSFORM STREAM TESTS
# ==========================================


def test_infracoes_transform_missing_dat_limt_recu():
    df = pd.DataFrame(
        {
            "NUM_AI": ["12345-A"],
            "DAT_LIMT_RECU": [""],
        }
    )
    stream = InfracoesTransformStream(
        datetime_format="%d/%m/%Y", date_format="%d/%m/%Y"
    )

    with pytest.raises(InvalidDocumentDataError) as exc_info:
        stream.transform(df)
    assert "DAT_LIMT_RECU" in str(exc_info.value)


def test_infracoes_transform_skips_blank_rows():
    df = pd.DataFrame(
        {
            "NUM_AI": ["12345-A", "", None, "  ", "67890-B"],
            "DAT_LIMT_RECU": ["15/09/2026", "", None, "  ", "16/09/2026"],
            "NOM_CONC": ["Consórcio A", "", None, "  ", "Consórcio B"],
        }
    )
    stream = InfracoesTransformStream(
        datetime_format="%d/%m/%Y %H:%M", date_format="%d/%m/%Y"
    )
    records = stream.transform(df)
    assert len(records) == 2
    assert records[0]["NUM_AI"] == "12345-A"
    assert records[1]["NUM_AI"] == "67890-B"


def test_infracoes_transform_skips_duplicate_headers():
    df = pd.DataFrame(
        {
            "NUM_AI": ["12345-A", "NUM_AI", " num_ai ", "67890-B"],
            "DAT_LIMT_RECU": [
                "15/09/2026",
                "DAT_LIMT_RECU",
                "DAT_LIMT_RECU",
                "16/09/2026",
            ],
            "NOM_CONC": ["Consórcio A", "NOM_CONC", "NOM_CONC", "Consórcio B"],
        }
    )
    stream = InfracoesTransformStream(
        datetime_format="%d/%m/%Y %H:%M", date_format="%d/%m/%Y"
    )
    records = stream.transform(df)
    assert len(records) == 2
    assert records[0]["NUM_AI"] == "12345-A"
    assert records[1]["NUM_AI"] == "67890-B"


def test_infracoes_transform_all_rows_blank_returns_empty():
    df = pd.DataFrame(
        {
            "NUM_AI": ["", None, "NUM_AI"],
            "DAT_LIMT_RECU": ["", None, "DAT_LIMT_RECU"],
        }
    )
    stream = InfracoesTransformStream(
        datetime_format="%d/%m/%Y %H:%M", date_format="%d/%m/%Y"
    )
    records = stream.transform(df)
    assert records == []


def test_infracoes_transform_invalid_date_meaningful_error():
    df = pd.DataFrame(
        {
            "NUM_AI": ["12345-A", "67890-B"],
            "DAT_LIMT_RECU": [
                "15/09/2026",
                "2026-09-16",
            ],  # Line 3 has ISO date instead of dd/mm/YYYY
        }
    )
    stream = InfracoesTransformStream(
        datetime_format="%d/%m/%Y %H:%M", date_format="%d/%m/%Y"
    )
    with pytest.raises(InvalidDocumentDataError) as exc_info:
        stream.transform(df)

    msg = str(exc_info.value)
    assert "DAT_LIMT_RECU" in msg
    assert "linha 3" in msg
    assert "2026-09-16" in msg
    assert "%d/%m/%Y" in msg


def test_infracoes_transform_cod_linh_formatting():
    df = pd.DataFrame(
        {
            "NUM_AI": ["12345-A", "67890-B", "99999-C"],
            "COD_LINH": [61.0, "6101", np.nan],
            "NOM_LINH": ["Linha 61", "Linha 6101", np.nan],
            "DAT_LIMT_RECU": ["15/09/2026", "16/09/2026", "17/09/2026"],
        }
    )
    stream = InfracoesTransformStream(
        datetime_format="%d/%m/%Y %H:%M", date_format="%d/%m/%Y"
    )
    records = stream.transform(df)
    assert len(records) == 3
    assert records[0]["COD_LINH"] == "61"
    assert records[0]["NOM_LINH"] == "Linha 61"
    assert records[1]["COD_LINH"] == "6101"
    assert records[2]["COD_LINH"] == ""
    assert records[2]["NOM_LINH"] == ""


def test_infracoes_transform_drops_empty_unnamed_columns():
    df = pd.DataFrame(
        {
            "NUM_AI": ["12345-A", "67890-B"],
            "DAT_LIMT_RECU": ["15/09/2026", "16/09/2026"],
            "Unnamed: 2": [None, None],
            "Unnamed: 3": ["", "  "],
            "Unnamed: 4": [np.nan, "valid_extra"],
        }
    )
    stream = InfracoesTransformStream(
        datetime_format="%d/%m/%Y %H:%M", date_format="%d/%m/%Y"
    )
    records = stream.transform(df)
    assert len(records) == 2
    assert "Unnamed: 2" not in records[0]
    assert "Unnamed: 3" not in records[0]
    # Unnamed: 4 was not completely empty, so it was preserved
    assert "Unnamed: 4" in records[0]
    assert records[1]["Unnamed: 4"] == "valid_extra"


# ==========================================
# DOCX INPUT STREAM TESTS
# ==========================================


def test_docx_stream_empty_num_recurso():
    doc = Document()
    table = doc.add_table(rows=2, cols=4)
    table.cell(0, 0).text = "RECURSO"
    table.cell(0, 1).text = "AUTO DE INFRAÇÃO"
    table.cell(0, 2).text = "RECORRENTE"
    table.cell(0, 3).text = "DECISÃO"

    # Row with empty recurso number
    table.cell(1, 0).text = ""
    table.cell(1, 1).text = "12345A"
    table.cell(1, 2).text = "Consorcio"
    table.cell(1, 3).text = "PROCEDENTE"

    stream = RecursosDocxInputStream(first_instance=True)
    with pytest.raises(InvalidDocumentDataError) as exc_info:
        list(stream.process_table(table, "2026-01-01", 1))
    assert "número do recurso está vazio" in str(exc_info.value)


def _create_sample_doc_stream(doc: DocxDocument) -> BytesIO:
    stream = BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream


def _populate_test_table(
    table: Table, rows_data: list[tuple[str, str, str, str]]
) -> None:
    table.cell(0, 0).text = "RECURSO"
    table.cell(0, 1).text = "AUTO DE INFRAÇÃO"
    table.cell(0, 2).text = "RECORRENTE"
    table.cell(0, 3).text = "DECISÃO"
    for row_idx, data in enumerate(rows_data, start=1):
        for col_idx, val in enumerate(data):
            table.cell(row_idx, col_idx).text = val


def test_iter_block_items_preserves_order():
    doc = Document()
    doc.add_paragraph("Paragraph 1")
    doc.add_table(rows=1, cols=1)
    doc.add_paragraph("Paragraph 2")
    doc.add_table(rows=1, cols=1)
    doc.add_table(rows=1, cols=1)

    parser = RecursosDocxInputStream()
    items = list(parser._iter_block_items(doc))
    assert len(items) == 5

    assert isinstance(items[0], Paragraph)
    assert items[0].text == "Paragraph 1"
    assert isinstance(items[1], Table)
    assert isinstance(items[2], Paragraph)
    assert items[2].text == "Paragraph 2"
    assert isinstance(items[3], Table)
    assert isinstance(items[4], Table)


def test_docx_stream_one_ata_multiple_tables():
    doc = Document()
    doc.add_paragraph(
        "PUBLICADO NO DIARIO OFICIAL DO MUNICIPIO DE BELO HORIZONTE EM 15/05/2026"
    )
    doc.add_paragraph("ATA DA 5ª SESSÃO ORDINÁRIA")

    t1 = doc.add_table(rows=2, cols=4)
    _populate_test_table(t1, [("101/2026", "11111A", "Consórcio 1", "IMPROCEDENTE")])

    t2 = doc.add_table(rows=2, cols=4)
    _populate_test_table(t2, [("102/2026", "22222A", "Consórcio 2", "PROCEDENTE")])

    t3 = doc.add_table(rows=2, cols=4)
    _populate_test_table(t3, [("103/2026", "33333A", "Consórcio 3", "PROCEDENTE")])

    stream = RecursosDocxInputStream(first_instance=True)
    results = list(stream.read(_create_sample_doc_stream(doc)))

    assert len(results) == 3
    assert results[0]["NUM_ATA"] == "5"
    assert results[0]["NUM_RECURSO"] == "101/2026"
    assert results[1]["NUM_ATA"] == "5"
    assert results[1]["NUM_RECURSO"] == "102/2026"
    assert results[2]["NUM_ATA"] == "5"
    assert results[2]["NUM_RECURSO"] == "103/2026"


def test_docx_stream_multiple_atas_multiple_tables():
    doc = Document()
    doc.add_paragraph(
        "PUBLICADO NO DIARIO OFICIAL DO MUNICIPIO DE BELO HORIZONTE EM 15/05/2026"
    )
    doc.add_paragraph("ATA DA 10ª SESSÃO")

    t1 = doc.add_table(rows=2, cols=4)
    _populate_test_table(t1, [("101/2026", "11111A", "Consórcio 1", "IMPROCEDENTE")])

    t2 = doc.add_table(rows=2, cols=4)
    _populate_test_table(t2, [("102/2026", "22222A", "Consórcio 2", "PROCEDENTE")])

    doc.add_paragraph("ATA DA 11ª SESSÃO")

    t3 = doc.add_table(rows=2, cols=4)
    _populate_test_table(t3, [("103/2026", "33333A", "Consórcio 3", "PROCEDENTE")])

    t4 = doc.add_table(rows=2, cols=4)
    _populate_test_table(t4, [("104/2026", "44444A", "Consórcio 4", "IMPROCEDENTE")])

    stream = RecursosDocxInputStream(first_instance=True)
    results = list(stream.read(_create_sample_doc_stream(doc)))

    assert len(results) == 4
    assert results[0]["NUM_ATA"] == "10"
    assert results[0]["NUM_RECURSO"] == "101/2026"
    assert results[1]["NUM_ATA"] == "10"
    assert results[1]["NUM_RECURSO"] == "102/2026"
    assert results[2]["NUM_ATA"] == "11"
    assert results[2]["NUM_RECURSO"] == "103/2026"
    assert results[3]["NUM_ATA"] == "11"
    assert results[3]["NUM_RECURSO"] == "104/2026"


def test_docx_stream_segunda_instancia_with_atas_raises_incorrect_instance():
    doc = Document()
    doc.add_paragraph(
        "PUBLICADO NO DIARIO OFICIAL DO MUNICIPIO DE BELO HORIZONTE EM 01/01/2026"
    )

    doc.add_paragraph("ATA DA 4ª SESSÃO ORDINÁRIA")

    t0 = doc.add_table(2, 4)
    _populate_test_table(t0, [("101/2026", "11111A", "Consórcio 1", "IMPROCEDENTE")])

    stream = RecursosDocxInputStream(first_instance=False)
    with pytest.raises(IncorrectInstanceError):
        list(stream.read(_create_sample_doc_stream(doc)))


def test_docx_stream_table_before_initial_ata_raises_error():
    doc = Document()
    doc.add_paragraph(
        "PUBLICADO NO DIARIO OFICIAL DO MUNICIPIO DE BELO HORIZONTE EM 15/05/2026"
    )

    # Table appears BEFORE any Ata header
    t1 = doc.add_table(rows=2, cols=4)
    _populate_test_table(t1, [("101/2026", "11111A", "Consórcio 1", "IMPROCEDENTE")])

    doc.add_paragraph("ATA DA 5ª SESSÃO ORDINÁRIA")
    t2 = doc.add_table(rows=2, cols=4)
    _populate_test_table(t2, [("102/2026", "22222A", "Consórcio 2", "PROCEDENTE")])

    stream = RecursosDocxInputStream(first_instance=True)
    with pytest.raises(QuantityOfAtasMismatchError) as exc_info:
        list(stream.read(_create_sample_doc_stream(doc)))
    assert "Tabela encontrada antes de qualquer ata" in str(exc_info.value)


def test_docx_stream_no_atas_raises_error():
    doc = Document()
    doc.add_paragraph(
        "PUBLICADO NO DIARIO OFICIAL DO MUNICIPIO DE BELO HORIZONTE EM 15/05/2026"
    )
    t1 = doc.add_table(rows=2, cols=4)
    _populate_test_table(t1, [("101/2026", "11111A", "Consórcio 1", "IMPROCEDENTE")])

    stream = RecursosDocxInputStream(first_instance=True)
    with pytest.raises(QuantityOfAtasMismatchError):
        list(stream.read(_create_sample_doc_stream(doc)))


def test_docx_stream_segunda_instancia_multiple_tables():
    doc = Document()
    doc.add_paragraph(
        "PUBLICADO NO DIARIO OFICIAL DO MUNICIPIO DE BELO HORIZONTE EM 15/05/2026"
    )
    t1 = doc.add_table(rows=2, cols=4)
    _populate_test_table(t1, [("201/2026", "11111A", "Consórcio 1", "IMPROCEDENTE")])
    t2 = doc.add_table(rows=2, cols=4)
    _populate_test_table(t2, [("202/2026", "22222A", "Consórcio 2", "PROCEDENTE")])

    stream = RecursosDocxInputStream(first_instance=False)
    results = list(stream.read(_create_sample_doc_stream(doc)))

    assert len(results) == 2
    assert "NUM_ATA" not in results[0]
    assert "NUM_ATA" not in results[1]
    assert results[0]["NUM_RECURSO"] == "201/2026"
    assert results[1]["NUM_RECURSO"] == "202/2026"


def test_docx_stream_dat_publ_invalid_format():
    doc = Document()
    doc.add_paragraph(
        "PUBLICADO NO DIARIO OFICIAL DO MUNICIPIO DE BELO HORIZONTE EM 2026/13/01"
    )

    t1 = doc.add_table(rows=2, cols=4)
    _populate_test_table(t1, [("201/2026", "11111A", "Consórcio 1", "IMPROCEDENTE")])
    t2 = doc.add_table(rows=2, cols=4)
    _populate_test_table(t2, [("202/2026", "22222A", "Consórcio 2", "PROCEDENTE")])

    stream = RecursosDocxInputStream(first_instance=False)
    with pytest.raises(PublicationDateNotFoundError):
        list(stream.read(_create_sample_doc_stream(doc)))


def test_docx_stream_dat_publ_invalid_format_II():
    doc = Document()
    doc.add_paragraph(
        "PUBLICADO NO DIARIO OFICIAL DO MUNICIPIO DE BELO HORIZONTE EM 00/0000/0000"
    )

    stream = RecursosDocxInputStream()
    with pytest.raises(PublicationDateNotFoundError):
        list(stream.read(_create_sample_doc_stream(doc)))
